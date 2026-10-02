#!/usr/bin/env python3
"""Build the Austrian district (Bezirke, 94) map layer.

Results: BMI Nationalratswahl 2024 final result per Gemeinde (official
xlsx, 16 Oct 2024). Unit = Bezirk: row G<ddd>00 (Gemeinden + the Bezirk's
postal votes). The regional-constituency postal rows G<d>X099 are counted
outside the Bezirke and are distributed proportionally within their RWK so
both layers sum to the national result. Vienna = row G90000 (whole city).
Geometry: geoBoundaries AUT ADM2 (94 features, Wien as one).

Writes img/austria_bezirke.svg and patches the austria block in
js/config.js with a map2 layer (and official per-Bundesland gebiete).

Usage: python scraper/build_austria_bezirke.py [--force]
"""
import json
import os
import re
import sys
import urllib.request

sys.path.insert(0, os.path.dirname(__file__))
import build_map_svg as bm

ROOT = os.path.join(os.path.dirname(__file__), "..")
XLSX = os.path.join(ROOT, "scraper", ".cache", "at_nrw2024.xlsx")
XLSX_URL = ("https://www.bmi.gv.at/412/nationalratswahlen/"
            "nationalratswahl_2024/files/"
            "endgueltiges_ergebnis_beschluss_bundeswahlbehoerde_16102024.xlsx")
OUT_SVG = os.path.join(ROOT, "img", "austria_bezirke.svg")
BMV_SVG = os.path.join(ROOT, "bmv", "img", "austria_bezirke.svg")
COMBINED = os.path.join(ROOT, "scraper", ".cache", "at_bezirke.geojson")
ATTRIBUTION = ("Ergebnisse: BMI (NRW 2024); Geometrie: geoBoundaries.org "
               "(gbOpen)")

PARTY_COL = {"oevp": 6, "spoe": 8, "fpoe": 10, "gruene": 12, "neos": 14,
             "kpoe": 26}
LANDS = {"1": "burgenland", "2": "carinthia", "3": "lower_austria",
         "4": "upper_austria", "5": "salzburg", "6": "styria",
         "7": "tyrol", "8": "vorarlberg", "9": "vienna"}
NAME_EXC = {"wien": "wien_stadt",
            "krems_an_der_donau": "krems_an_der_donau_stadt"}


def norm(name):
    s = name.replace("ß", "ss")
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue")):
        s = s.replace(a, b)
    s = s.lower().replace("(", " ").replace(")", " ")
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def main():
    force = "--force" in sys.argv
    if force or not os.path.isfile(XLSX):
        req = urllib.request.Request(XLSX_URL,
                                     headers={"User-Agent": "Mozilla/5.0"})
        data = urllib.request.urlopen(req, timeout=180).read()
        with open(XLSX, "wb") as fh:
            fh.write(data)
        print("downloaded xlsx:", len(data), "bytes")

    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    meig = {str(r[0]): r for r in rows[1:] if r[0]}

    def votes(gkz, party):
        return meig[gkz][PARTY_COL[party]] or 0

    def valid(gkz):
        return meig[gkz][5] or 0

    votes_b = {}
    for gkz in meig:
        if re.fullmatch(r"G[1-8]\d\d00", gkz) and not re.fullmatch(r"G\d0000",
                                                                   gkz):
            votes_b[gkz] = {p: votes(gkz, p) for p in PARTY_COL}
            votes_b[gkz]["_v"] = valid(gkz)
    votes_b["G90000"] = {p: votes("G90000", p) for p in PARTY_COL}
    votes_b["G90000"]["_v"] = valid("G90000")
    print("bezirke:", len(votes_b))

    # distribute the RWK-level postal rows proportionally by valid votes
    for d in range(1, 9):
        bez = [k for k in votes_b if re.fullmatch("G%d\\d\\d00" % d, k)]
        for rk in sorted(k for k in meig if re.fullmatch(
                "G%d[A-G]000" % d, k)):
            wkk = re.sub(r"00$", "99", rk)
            extra = valid(wkk)
            if not extra:
                continue
            target = valid(rk) - extra
            items = sorted(((k, votes_b[k]["_v"]) for k in bez),
                           key=lambda x: -x[1])

            def dfs(i, rem):
                if rem == 0:
                    return []
                if i >= len(items) or rem < 0:
                    return None
                for j in range(i, len(items)):
                    k, val = items[j]
                    got = dfs(j + 1, rem - val)
                    if got is not None:
                        return [k] + got
                return None

            members = dfs(0, target)
            assert members, f"no RWK match for {meig[rk][1]} target {target}"
            for k in members:
                w = votes_b[k]["_v"] / target
                votes_b[k]["_v"] += extra * w
                for p in PARTY_COL:
                    votes_b[k][p] += votes(wkk, p) * w

    gebiete, names = {}, {}
    for gkz, acc in votes_b.items():
        if gkz == "G90000":
            key, disp = "wien_stadt", "Wien"
        else:
            raw = meig[gkz][1]
            key = NAME_EXC.get(norm(raw), norm(raw))
            disp = raw
        gebiete[key] = {p: round(acc[p] * 100 / acc["_v"], 2)
                        for p in PARTY_COL}
        names[key] = disp
    print("units with shares:", len(gebiete))

    land_gebiete = {}
    for d, lk in LANDS.items():
        row = "G%s0000" % d
        v = valid(row)
        land_gebiete[lk] = {p: round(votes(row, p) * 100 / v, 2)
                            for p in PARTY_COL}

    gj = bm.load_geojson(
        "https://www.geoboundaries.org/api/current/gbOpen/AUT/ADM2/")
    features, matched = [], set()
    for f in gj["features"]:
        key = norm(f["properties"]["shapeName"])
        if key not in gebiete:
            print("  no results for", f["properties"]["shapeName"], "->", key)
            continue
        matched.add(key)
        features.append({"type": "Feature", "properties": {"bezirk": key},
                         "geometry": f["geometry"]})
    print("features:", len(features), "| results without geometry:",
          sorted(set(gebiete) - matched))
    assert len(features) == 94, "expected 94 bezirke"
    with open(COMBINED, "w", encoding="utf8") as fh:
        json.dump({"type": "FeatureCollection", "features": features}, fh)
    sys.argv = ["build_map_svg.py", "--geojson", COMBINED,
                "--name-field", "bezirk", "--attr", "id", "--no-prefix",
                "--out", OUT_SVG, "--attribution", ATTRIBUTION, "--force"]
    bm.main()
    import shutil
    shutil.copy2(OUT_SVG, BMV_SVG)

    # verify: land = valid-weighted sum of its bezirke
    for d, lk in LANDS.items():
        if d == "9":
            continue
        bez = [k for k in votes_b if re.fullmatch("G%s\\d\\d00" % d, k)]
        tv = sum(votes_b[k]["_v"] for k in bez)
        worst = max(PARTY_COL, key=lambda p: abs(
            sum(votes_b[k][p] for k in bez) * 100 / tv - land_gebiete[lk][p]))
        diff = sum(votes_b[k][worst] for k in bez) * 100 / tv - \
            land_gebiete[lk][worst]
        print(f"  {lk:15s} sum-check worst {worst} {diff:+.2f}pp")

    cfg_path = os.path.join(ROOT, "js", "config.js")
    text = open(cfg_path, encoding="utf8").read()
    start = text.index("\n  austria: {")
    m = re.search(r"\n  \w+: \{", text[start + 1:])
    end = start + 1 + m.start()
    block = text[start:end]

    # layer 1: official per-Bundesland gebiete
    gm = re.search(r"\n(\s+)gebiete: \{", block)
    indent = gm.group(1)
    depth, k = 0, gm.end() - 1
    while k < len(block):
        if block[k] == "{":
            depth += 1
        elif block[k] == "}":
            depth -= 1
            if depth == 0:
                break
        k += 1
    new_body = json.dumps(land_gebiete, ensure_ascii=False, indent=len(indent))
    block = (block[:gm.start()] + "\n" + indent + "gebiete: " +
             new_body + block[k + 1:])

    map2 = {"svg": "img/austria_bezirke.svg", "selector": "id",
            "districts": {k2: k2 for k2 in gebiete},
            "gebiete": gebiete, "names": names, "label": "Bezirke (94)"}
    mm = re.search(r"\n\s+map: \{", block)
    depth, k = 0, mm.end() - 1
    while k < len(block):
        if block[k] == "{":
            depth += 1
        elif block[k] == "}":
            depth -= 1
            if depth == 0:
                break
        k += 1
    m2 = re.search(r"\n\s+map2: \{", block)
    if m2:
        depth, k2 = 0, m2.end() - 1
        while k2 < len(block):
            if block[k2] == "{":
                depth += 1
            elif block[k2] == "}":
                depth -= 1
                if depth == 0:
                    break
            k2 += 1
        block = block[:m2.start()] + "\n      map2: " + \
            json.dumps(map2, ensure_ascii=False, indent=6) + block[k2 + 1:]
    else:
        block = block[:k + 1] + ",\n      map2: " + \
            json.dumps(map2, ensure_ascii=False, indent=6) + block[k + 1:]
    open(cfg_path, "w", encoding="utf8").write(text[:start] + block + text[end:])
    print("patched config.js (austria map2 + official land gebiete)")


if __name__ == "__main__":
    main()
